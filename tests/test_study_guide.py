"""
Unit and Integration Tests for Section-Based Map-Reduce Study Guide Pipeline.
"""

import io
import pytest
import docx
from fastapi.testclient import TestClient
from unittest.mock import patch, MagicMock
from types import SimpleNamespace

from generation.study_guide import (
    partition_chunks_into_sections,
    build_batch_packed_prompt,
    _call_extractive_study_guide,
    generate_comprehensive_study_guide,
)
from generation.docx_exporter import build_study_guide_docx
from api.main import app

@pytest.fixture(autouse=True)
def mock_env_settings():
    mock_settings = SimpleNamespace(neo4j_uri="", gemini_api_key="")
    with patch("api.main.get_settings", return_value=mock_settings):
        yield

client = TestClient(app)


# ---------------------------------------------------------------------------
# Unit Tests: Partitioning & Outline Formation
# ---------------------------------------------------------------------------

def test_partition_chunks_zero_loss():
    """Verify that every single chunk is retained in the section partition."""
    chunks = [
        {
            "chunk_id": "c1",
            "text": "Introduction text describing the motivation.",
            "metadata": {"doc_id": "Paper.pdf", "chunk_index": 0, "page_number": 1, "section_header": "1. Introduction"}
        },
        {
            "chunk_id": "c2",
            "text": "Background details and related works.",
            "metadata": {"doc_id": "Paper.pdf", "chunk_index": 1, "page_number": 2, "section_header": "1. Introduction"}
        },
        {
            "chunk_id": "c3",
            "text": "Proposed GraphRAG hybrid retrieval architecture.",
            "metadata": {"doc_id": "Paper.pdf", "chunk_index": 2, "page_number": 3, "section_header": "2. Architecture"}
        },
        {
            "chunk_id": "c4",
            "text": "Experimental setup and benchmark metrics.",
            "metadata": {"doc_id": "Paper.pdf", "chunk_index": 3, "page_number": 4, "section_header": "3. Experiments"}
        },
        {
            "chunk_id": "c5",
            "text": "Conclusion and future research directions.",
            "metadata": {"doc_id": "Paper.pdf", "chunk_index": 4, "page_number": 5, "section_header": "4. Conclusion"}
        }
    ]

    sections = partition_chunks_into_sections(chunks)

    # 4 distinct section headers
    assert len(sections) == 4
    assert sections[0]["section_title"] == "1. Introduction"
    assert len(sections[0]["chunks"]) == 2
    assert sections[1]["section_title"] == "2. Architecture"
    assert sections[2]["section_title"] == "3. Experiments"
    assert sections[3]["section_title"] == "4. Conclusion"

    # Verify total chunks preserved across all sections
    total_chunks = sum(len(s["chunks"]) for s in sections)
    assert total_chunks == 5


def test_partition_chunks_without_headers():
    """Verify that chunks with empty headers are clustered cleanly by page."""
    chunks = [
        {"chunk_id": "c1", "text": "Page 1 intro", "metadata": {"doc_id": "Doc.pdf", "chunk_index": 0, "page_number": 1}},
        {"chunk_id": "c2", "text": "Page 1 detail", "metadata": {"doc_id": "Doc.pdf", "chunk_index": 1, "page_number": 1}},
        {"chunk_id": "c3", "text": "Page 2 analysis", "metadata": {"doc_id": "Doc.pdf", "chunk_index": 2, "page_number": 2}},
        {"chunk_id": "c4", "text": "Page 2 findings", "metadata": {"doc_id": "Doc.pdf", "chunk_index": 3, "page_number": 2}},
    ]
    sections = partition_chunks_into_sections(chunks)
    assert len(sections) >= 1
    total = sum(len(s["chunks"]) for s in sections)
    assert total == 4


def test_build_batch_packed_prompt():
    """Verify prompt formatting, citation numbering, and coverage rule."""
    chunks = [
        {"chunk_id": "c1", "text": "Self-attention mechanism", "metadata": {"doc_id": "Attention.pdf", "chunk_index": 0, "page_number": 2, "section_header": "Attention"}}
    ]
    sections = partition_chunks_into_sections(chunks)
    prompt, citations = build_batch_packed_prompt(sections, "Transformer Guide")

    assert "CRITICAL COVERAGE RULE" in prompt
    assert "TOTAL SECTIONS COVERED: 1" in prompt
    assert "Self-attention mechanism" in prompt
    assert len(citations) == 1
    assert citations[0]["citation_number"] == 1


# ---------------------------------------------------------------------------
# Unit Tests: Fallback Cascades & Word Export
# ---------------------------------------------------------------------------

def test_deterministic_extractive_fallback():
    """Verify that the extractive fallback produces a clean 5-part guide without any LLM."""
    chunks = [
        {"chunk_id": "c1", "text": "Deep learning models require optimization. Gradient descent converges iteratively.", "metadata": {"doc_id": "Notes.pdf", "chunk_index": 0, "page_number": 1, "section_header": "Optimization"}},
        {"chunk_id": "c2", "text": "Learning rate controls step size. Large learning rates risk divergence.", "metadata": {"doc_id": "Notes.pdf", "chunk_index": 1, "page_number": 2, "section_header": "Learning Rate"}},
    ]
    sections = partition_chunks_into_sections(chunks)
    markdown = _call_extractive_study_guide(sections, "Machine Learning")

    assert "# Executive Briefing" in markdown
    assert "Optimization" in markdown
    assert "Learning Rate" in markdown
    assert "Active Recall Review Questions" in markdown

    # Verify rendering into docx
    buf = build_study_guide_docx(
        notebook_name="Machine Learning",
        citations=[{"citation_number": 1, "doc_id": "Notes.pdf", "page_number": 1, "section_header": "Optimization", "text_preview": "Deep learning..."}],
        markdown_content=markdown
    )
    assert buf.getvalue().startswith(b"PK\x03\x04")

    # Read back docx
    doc = docx.Document(buf)
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "Optimization" in all_text
    assert "Learning Rate" in all_text


def test_markdown_table_rendering_in_docx():
    """Verify that markdown tables (like Glossary) render as Word tables."""
    table_markdown = """# Glossary of Essential Terminology
| Term | Definition | Primary Source |
| :--- | :--- | :--- |
| Attention | Mechanism allowing dynamic focus | Attention.pdf |
| ChromaDB | Serverless vector database | Architecture.docx |
"""
    buf = build_study_guide_docx(
        notebook_name="AI Terms",
        markdown_content=table_markdown,
        citations=[]
    )
    doc = docx.Document(buf)
    # 1 callout box table + 1 markdown table = at least 2 tables
    assert len(doc.tables) >= 2
    # Find glossary table
    tbl = doc.tables[-1]
    header_cells = [c.text for c in tbl.rows[0].cells]
    assert "Term" in header_cells
    assert "Definition" in header_cells


# ---------------------------------------------------------------------------
# Integration Tests: FastAPI Endpoint
# ---------------------------------------------------------------------------

def test_api_export_study_guide_endpoint():
    """Verify POST /notebooks/{id}/export/study-guide returns a valid .docx."""
    nb_res = client.post("/notebooks", json={"name": "End to End Study Guide"})
    assert nb_res.status_code == 200
    nb_id = nb_res.json()["notebook_id"]

    res = client.post(f"/notebooks/{nb_id}/export/study-guide")
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in res.headers["content-type"]
    assert res.content.startswith(b"PK\x03\x04")

    doc = docx.Document(io.BytesIO(res.content))
    assert len(doc.paragraphs) > 0
